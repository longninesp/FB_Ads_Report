#!/usr/bin/env python3
"""Convert a Meta Ads Manager bulk export into reusable JSON templates."""

from __future__ import annotations

import argparse
import hashlib
import json
import shutil
from collections import Counter
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


def clean_id(value: Any) -> str:
    text = str(value or "")
    return text.split(":", 1)[-1] if ":" in text else text


def non_empty_record(headers: list[str], values: tuple[Any, ...]) -> dict[str, Any]:
    return {
        header: value
        for header, value in zip(headers, values)
        if header and value not in (None, "")
    }


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def write_json(path: Path, payload: Any) -> None:
    path.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Archive and normalize a Meta Ads Manager XLSX bulk export."
    )
    parser.add_argument("xlsx", type=Path)
    parser.add_argument("output_dir", type=Path)
    args = parser.parse_args()

    source = args.xlsx.expanduser().resolve()
    output_dir = args.output_dir.expanduser().resolve()
    output_dir.mkdir(parents=True, exist_ok=True)

    workbook = load_workbook(source, read_only=False, data_only=True)
    worksheet = workbook.active
    rows = list(worksheet.iter_rows(values_only=True))
    if len(rows) < 2:
        raise SystemExit("The workbook has no campaign rows.")

    headers = [str(value or "").strip() for value in rows[0]]
    if len(headers) != len(set(headers)):
        duplicates = sorted(
            header for header, count in Counter(headers).items() if header and count > 1
        )
        raise SystemExit(f"Duplicate headers are not supported: {duplicates}")

    campaign_end = headers.index("Ad Set ID")
    adset_end = headers.index("Story ID")
    data_rows = [row for row in rows[1:] if any(value not in (None, "") for value in row)]

    campaign_headers = headers[:campaign_end]
    adset_headers = headers[campaign_end:adset_end]
    ad_headers = headers[adset_end:]

    campaigns: dict[str, dict[str, Any]] = {}
    adsets: dict[str, dict[str, Any]] = {}
    ads: list[dict[str, Any]] = []
    raw_rows: list[dict[str, Any]] = []

    for row in data_rows:
        raw_rows.append(non_empty_record(headers, row))

        campaign_values = row[:campaign_end]
        campaign = non_empty_record(campaign_headers, campaign_values)
        campaign_id = clean_id(campaign.get("Campaign ID"))
        campaigns.setdefault(campaign_id, campaign)

        adset_values = row[campaign_end:adset_end]
        adset = non_empty_record(adset_headers, adset_values)
        adset_id = clean_id(adset.get("Ad Set ID"))
        adsets.setdefault(
            adset_id,
            {
                "campaign_id": campaign_id,
                "raw_config": adset,
            },
        )

        ad_values = row[adset_end:]
        ads.append(non_empty_record(ad_headers, ad_values))

    creative_inventory = []
    for ad in ads:
        creative_inventory.append(
            {
                "ad_id": clean_id(ad.get("Ad ID")),
                "status": ad.get("Ad Status"),
                "ad_name": ad.get("Ad Name"),
                "creative_type": ad.get("Creative Type"),
                "title": ad.get("Title"),
                "body": ad.get("Body"),
                "call_to_action": ad.get("Call to Action"),
                "image_hash": ad.get("Image Hash"),
                "image_file_name": ad.get("Image File Name"),
                "video_id": clean_id(ad.get("Video ID")),
                "video_file_name": ad.get("Video File Name"),
                "preview_link": ad.get("Preview Link"),
                "instagram_preview_link": ad.get("Instagram Preview Link"),
                "permalink": ad.get("Permalink"),
                "raw_config": ad,
            }
        )

    account_ids = sorted(
        {
            str(item["image_hash"]).split(":", 1)[0]
            for item in creative_inventory
            if item.get("image_hash") and ":" in str(item["image_hash"])
        }
    )
    template = {
        "source": {
            "file": source.name,
            "sha256": sha256(source),
            "sheet": worksheet.title,
            "columns": len(headers),
            "rows": len(data_rows),
            "inferred_ad_account_ids": account_ids,
        },
        "summary": {
            "campaigns": len(campaigns),
            "ad_sets": len(adsets),
            "ads": len(ads),
            "active_ads": sum(ad.get("Ad Status") == "ACTIVE" for ad in ads),
            "paused_ads": sum(ad.get("Ad Status") == "PAUSED" for ad in ads),
            "creative_types": dict(Counter(ad.get("Creative Type", "") for ad in ads)),
        },
        "campaigns": [
            {
                "campaign_id": campaign_id,
                "raw_config": config,
            }
            for campaign_id, config in campaigns.items()
        ],
        "ad_sets": [
            {"ad_set_id": adset_id, **config}
            for adset_id, config in adsets.items()
        ],
        "creation_notes": {
            "default_creation_status": "PAUSED",
            "currency_amounts_require_account_currency_confirmation": True,
            "placements_blank_means_preserve_or_confirm_automatic_placements": True,
            "asset_files_embedded_in_source": False,
            "reuse_existing_asset_ids_requires_access_to_source_ad_account": True,
        },
    }

    archived_source = output_dir / source.name
    shutil.copy2(source, archived_source)
    write_json(output_dir / "campaign_template.json", template)
    write_json(output_dir / "creative_inventory.json", creative_inventory)
    write_json(output_dir / "raw_non_empty_rows.json", raw_rows)

    print(
        json.dumps(
            {
                "output_dir": str(output_dir),
                **template["summary"],
                "source_sha256": template["source"]["sha256"],
            },
            ensure_ascii=False,
            indent=2,
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
